"""
The house style of the exports, shared by the stock cycle and competitor files.

exports.py keeps its helpers inside build_pptx and build_xlsx. These are the same helpers,
with the same fonts, colours, sizes and positions, lifted out so that every new file looks
like the company export and can also add its pages to it: Deck and Book each accept an
existing presentation or workbook and keep its page count running.
"""

import datetime
import io

from exports import BAND, GOLD, INK, LINE, NAVY, PALE, WHITE

UP, DOWN = '12703A', 'A5201A'
HEAD, BODY = 'Georgia', 'Calibri'
X0, CW = 0.6, 12.133


# ====================================================================== POWERPOINT =====
class Deck:
    def __init__(self, footer, prs=None, count=None):
        from pptx import Presentation
        from pptx.util import Inches
        if prs is None:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        self.prs, self.footer = prs, footer
        self.count = count if count is not None else {'n': 0}
        self.blank = prs.slide_layouts[6]

    @staticmethod
    def rgb(h):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(h)

    def text(self, slide, x, y, w, h, t, size=12, bold=False, color=INK, face=BODY, align=None,
             anchor=None, italic=False):
        from pptx.util import Inches, Pt
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor:
            tf.vertical_anchor = anchor
        lines = t if isinstance(t, list) else [t]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if isinstance(ln, tuple):
                ln, o = ln
            else:
                o = {}
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(o.get('size', size))
            r.font.bold = o.get('bold', bold)
            r.font.italic = italic
            r.font.name = o.get('face', face)
            r.font.color.rgb = self.rgb(o.get('color', color))
            if align:
                p.alignment = align
            if o.get('space'):
                p.space_before = Pt(o['space'])
        return tb

    def cover(self, title, lines):
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = self.rgb(NAVY)
        self.text(s, X0, 1.6, 11.8, 1.8, title, size=40, color=WHITE, face=HEAD,
                  anchor=MSO_ANCHOR.BOTTOM)
        rule = s.shapes.add_connector(1, Inches(X0), Inches(3.75), Inches(X0 + 1.4), Inches(3.75))
        rule.line.color.rgb = self.rgb(GOLD)
        rule.line.width = Pt(2)
        self.text(s, X0, 4.0, 11.8, 1.6,
                  [(t, {'color': WHITE if i == 0 else 'C9D3E3', 'size': 16 if i == 0 else 13,
                        'space': 6 if i else 0}) for i, t in enumerate(lines)])
        return s

    def content(self, title_text, sub=None):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
        self.count['n'] += 1
        s = self.prs.slides.add_slide(self.blank)
        self.text(s, X0, 0.38, CW, 0.75, title_text, size=24, color=NAVY, face=HEAD,
                  anchor=MSO_ANCHOR.BOTTOM)
        ln = s.shapes.add_connector(1, Inches(X0), Inches(1.22), Inches(X0 + CW), Inches(1.22))
        ln.line.color.rgb = self.rgb(LINE)
        ln.line.width = Pt(0.75)
        if sub:
            self.text(s, X0, 1.3, CW, 0.5, sub, size=11, color=NAVY)
        self.text(s, X0, 7.0, 9, 0.3, self.footer, size=9, color=NAVY)
        self.text(s, 12.1, 7.0, 0.63, 0.3, str(self.count['n'] + 1), size=9, color=NAVY,
                  align=PP_ALIGN.RIGHT)
        return s

    def table(self, slide, x, y, colw, rows, head=True, size=10.5, rowh=0.32, aligns=None,
              first_bold=True, col_fill=None, styles=None, group_rows=()):
        """The export table. col_fill tints whole columns (the company being analysed),
        styles maps (row, col) to {'color', 'bold'}, and group_rows are full-width labels."""
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
        nr, nc = len(rows), len(rows[0])
        shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)),
                                     Inches(rowh * nr))
        tbl = shp.table
        tblPr = shp._element.graphic.graphicData.tbl.tblPr
        tblPr.set('bandRow', '0')
        tblPr.set('firstRow', '0')
        style_id = tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
        if style_id is not None:
            style_id.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
        for j, w in enumerate(colw):
            tbl.columns[j].width = Inches(w)
        aligns = aligns or (['l'] + ['r'] * (nc - 1))
        amap = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}
        col_fill, styles = col_fill or {}, styles or {}
        two_line_head = head and any('\n' in str(v) for v in rows[0])
        band = 0
        for i in range(nr):
            tbl.rows[i].height = Inches(max(rowh, 0.55) if (i == 0 and two_line_head) else rowh)
            is_head = head and i == 0
            is_group = i in group_rows
            if not is_head and not is_group:
                band += 1
            for j in range(nc):
                cell = tbl.cell(i, j)
                val = rows[i][j]
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                o = styles.get((i, j), {})
                colour = (WHITE if is_head else NAVY if (is_group or (j == 0 and first_bold))
                          else INK)
                for k, part in enumerate(('' if val is None else str(val)).split('\n')):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = amap[aligns[j]] if not is_group else PP_ALIGN.LEFT
                    r = p.add_run()
                    r.text = part
                    r.font.name = BODY
                    # a second line in a header is the company name under its ticker
                    r.font.size = Pt((size - 1 if is_group else size) - (1.5 if k else 0))
                    r.font.bold = o.get('bold', (is_head and k == 0) or is_group or
                                        (first_bold and j == 0))
                    r.font.color.rgb = self.rgb(o.get('color', 'C9D3E3' if (is_head and k)
                                                      else colour))
                cell.fill.solid()
                if is_head:
                    cell.fill.fore_color.rgb = self.rgb(NAVY)
                elif is_group:
                    cell.fill.fore_color.rgb = self.rgb(PALE)
                elif j in col_fill:
                    cell.fill.fore_color.rgb = self.rgb(col_fill[j])
                else:
                    cell.fill.fore_color.rgb = self.rgb(BAND if band % 2 == 0 else WHITE)
            if is_group and nc > 1:
                tbl.cell(i, 0).merge(tbl.cell(i, nc - 1))
        return shp

    def line_chart(self, slide, x, y, w, h, dates, series, number_format='#,##0.00'):
        """A native line chart, styled like the one in the company export.

        series: (name, values, colour hex, width in points, dashed)."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt
        cd = CategoryChartData()
        cd.categories = list(dates)
        for name, values, *_ in series:
            cd.add_series(name, [None if v is None else round(float(v), 4) for v in values])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w),
                                    Inches(h), cd)
        ch = gf.chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
        ch.legend.font.name = BODY
        for ser, (_, _, colour, width, dashed) in zip(ch.series, series):
            ser.smooth = False
            ser.format.line.color.rgb = self.rgb(colour)
            ser.format.line.width = Pt(width)
            if dashed:
                ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ca, va = ch.category_axis, ch.value_axis
        skip = str(max(1, len(cd.categories) // 8))
        cax = ca._element
        anchor = cax.find(qn('c:noMultiLvlLbl'))
        for tag in ('c:tickLblSkip', 'c:tickMarkSkip'):
            el = cax.find(qn(tag))
            if el is None:
                el = cax.makeelement(qn(tag), {})
                if anchor is not None:
                    anchor.addprevious(el)
                else:
                    cax.append(el)
            el.set('val', skip)
        ca.tick_labels.font.size = Pt(9)
        ca.tick_labels.number_format = 'mmm yy'
        ca.tick_labels.number_format_is_linked = False
        va.tick_labels.font.size = Pt(9)
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = self.rgb(LINE)
        ca.format.line.color.rgb = self.rgb(LINE)
        va.format.line.color.rgb = self.rgb(LINE)
        return ch

    def save(self):
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


# ====================================================================== EXCEL ==========
class Book:
    def __init__(self, wb=None):
        from openpyxl import Workbook
        self.wb = wb if wb is not None else Workbook()
        self.fresh = wb is None

    @staticmethod
    def font(**k):
        from openpyxl.styles import Font
        base = dict(name='Arial', size=10, color=INK)
        base.update(k)
        return Font(**base)

    @staticmethod
    def fill(c):
        from openpyxl.styles import PatternFill
        return PatternFill('solid', fgColor=c)

    def sheet(self, title, widths=()):
        from openpyxl.utils import get_column_letter
        if self.fresh:
            ws, self.fresh = self.wb.active, False
        else:
            ws = self.wb.create_sheet()
        ws.title = title[:31]
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = 'MICS International, internal research'
        ws.oddFooter.right.text = 'Page &P of &N'
        return ws

    def title(self, ws, text, sub):
        ws['A1'] = text
        ws['A1'].font = self.font(size=18, bold=True, color=NAVY)
        ws['A2'] = sub
        ws['A2'].font = self.font(size=10, color=NAVY)
        return 4

    def heading(self, ws, r, text):
        ws.cell(r, 1, text).font = self.font(size=12, bold=True, color=NAVY)
        return r + 1

    def note(self, ws, r, text, cols, bold=False, colour=NAVY, chars=110):
        from openpyxl.styles import Alignment
        ws.cell(r, 1, text).font = self.font(color=colour, bold=bold)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 15 * max(1, len(text) // chars + 1)
        return r + 2

    def table(self, ws, r, head, rows, aligns=None, wrap_cols=(), formats=None,
              col_fill=None, styles=None, group_rows=()):
        """formats maps (row index, col index) or col index to a number format; group_rows are
        indexes into rows that are section labels across the table."""
        from openpyxl.styles import Alignment, Border, Side
        hair = Side(style='thin', color=LINE)
        gold = Side(style='medium', color=GOLD)
        aligns = aligns or (['left'] + ['right'] * (len(head) - 1))
        formats, col_fill, styles = formats or {}, col_fill or {}, styles or {}
        for j, h in enumerate(head, 1):
            c = ws.cell(r, j, h)
            c.font = self.font(bold=True, size=9, color=WHITE)
            c.fill = self.fill(NAVY)
            c.alignment = Alignment(horizontal=aligns[j - 1], vertical='center', wrap_text=True)
            c.border = Border(bottom=gold)
        ws.row_dimensions[r].height = 42 if any('\n' in str(h) for h in head) else 28
        r += 1
        band = 0
        for i, row in enumerate(rows):
            group = i in group_rows
            if not group:
                band += 1
            for j, val in enumerate(row, 1):
                c = ws.cell(r, j, val)
                o = styles.get((i, j - 1), {})
                c.font = self.font(bold=o.get('bold', group or j == 1),
                                   color=o.get('color', NAVY if (group or j == 1) else INK),
                                   size=9 if group else 10)
                c.alignment = Alignment(horizontal='left' if group else aligns[j - 1],
                                        vertical='top', wrap_text=(j - 1) in wrap_cols,
                                        indent=1 if (j > 1 and aligns[j - 1] == 'left') else 0)
                c.border = Border(bottom=hair)
                fmt = formats.get((i, j - 1), formats.get(j - 1))
                if fmt and isinstance(val, (int, float, datetime.date)):
                    c.number_format = fmt
                if group:
                    c.fill = self.fill(PALE)
                elif (j - 1) in col_fill:
                    c.fill = self.fill(col_fill[j - 1])
                elif band % 2 == 0:
                    c.fill = self.fill(BAND)
            if group and len(row) > 1:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(row))
            r += 1
        return r + 1

    def save(self):
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()
