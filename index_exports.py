"""
Excel and PowerPoint exports of an index analysis, in the same design as the company and
screener exports: tables first, one native line chart, a thin gold rule on the cover.

    build_index_xlsx(ctx) -> bytes
    build_index_pptx(ctx) -> bytes
"""

import io

import pandas as pd

import index_page as ip
import risk as rk
from exports import (BAND, GOLD, INK, LINE, NAVY, WHITE, fnum, plain)


def _rebased(ctx, years=5, freq='W-FRI'):
    a = ctx.get('primary')
    if a is None or len(a) < 30:
        return None
    a = a.loc[a.index[-1] - pd.DateOffset(years=years):].resample(freq).last().dropna()
    out = pd.DataFrame({'level': a / a.iloc[0] * 100})
    b = ctx.get('bench')
    if b is not None and not ctx.get('is_bench'):
        b = b.loc[a.index[0] - pd.Timedelta(days=40):a.index[-1]].resample(freq).last()
        b = b.reindex(out.index).ffill()
        if b.notna().sum() > 10:
            out['bench'] = b / b.dropna().iloc[0] * 100
    return out


def _returns_rows(ctx):
    e, fd = ctx['entry'], ctx['fund']
    ri, re_, rb = ctx.get('ret_idx') or {}, ctx.get('ret_etf') or {}, ctx.get('ret_bench') or {}
    head = ['Period']
    if ri:
        head.append(f'Index, {e.get("ccy")}, price')
    if re_:
        head.append(f'{e.get("etf")}, {fd.get("ccy") or ""}, with dividends')
    if not ctx.get('is_bench'):
        head.append('S&P 500, price')
    rows = []
    for lab, _ in rk.PERIODS:
        vals = ([ri.get(lab)] if ri else []) + ([re_.get(lab)] if re_ else []) + \
            ([] if ctx.get('is_bench') else [rb.get(lab)])
        if all(v is None for v in vals):
            continue
        name = {'YTD': 'Year to date'}.get(lab, lab) + (' a year' if lab in ('3Y', '5Y', '10Y') else '')
        rows.append([name] + vals)
    return head, rows


def _window_rows(ctx):
    out = []
    for lab in ctx.get('spans') or []:
        st = ctx['windows'].get(lab) or {}
        b, dd = st.get('bench') or {}, st.get('dd') or {}
        out.append([lab, st.get('change_pct'), b.get('bench_change_pct'), st.get('relative_pts'),
                    dd.get('worst_pct'), st.get('vol_pct'), b.get('beta'), b.get('corr')])
    return out


def _snapshot(ctx):
    e, fd = ctx['entry'], ctx['fund']
    main = ctx.get('ret_idx') or ctx.get('ret_etf') or {}
    w5 = (ctx['windows'].get('5Y') or {}).get('bench') or {}
    return [
        ['Level' if ctx['primary_is_index'] else f'{e.get("etf")} price',
         f'{fnum(ctx["last"])} {ctx.get("ccy") or ""}'.strip()],
        ['Year to date', fnum(main.get('YTD'), 'pct')], ['One year', fnum(main.get('1Y'), 'pct')],
        ['Five years, a year', fnum(main.get('5Y'), 'pct')],
        ['Below the record high', fnum((ctx.get('dd_full') or {}).get('now_below_high_pct'), 'pct')],
        ['Against the 200 day average', fnum(ctx.get('vs200'), 'pct')],
        ['Beta, five years', fnum(w5.get('beta'))], ['Correlation, five years', fnum(w5.get('corr'))],
        ['Fund', f'{fd.get("name") or e.get("etf") or "-"}'],
        ['Fund cost a year', '-' if fd.get('expense_pct') is None else f'{fd["expense_pct"]:.2f}%'],
        ['Fund P/E, trailing', fnum(fd.get('pe'), 'x')],
    ]


# ====================================================================== EXCEL ==========
def build_index_xlsx(ctx):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    e, fd = ctx['entry'], ctx['fund']
    as_of = ctx['end'].strftime('%d %B %Y').lstrip('0')
    font = lambda **k: Font(**dict(dict(name='Arial', size=10, color=INK), **k))
    fill = lambda c: PatternFill('solid', fgColor=c)
    hair, gold = Side(style='thin', color=LINE), Side(style='medium', color=GOLD)
    wb = Workbook()

    def sheet(title, first=False, widths=()):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.page_setup.orientation = 'landscape'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.oddFooter.left.text = 'MICS International, internal research'
        ws.oddFooter.right.text = 'Page &P of &N'
        return ws

    def title(ws, t, sub):
        ws['A1'] = t
        ws['A1'].font = font(size=18, bold=True, color=NAVY)
        ws['A2'] = sub
        ws['A2'].font = font(color=NAVY)
        return 4

    def table(ws, r, head, rows, fmt=None, left=(0,)):
        for j, h in enumerate(head, 1):
            c = ws.cell(r, j, h)
            c.font, c.fill = font(bold=True, size=9, color=WHITE), fill(NAVY)
            c.alignment = Alignment(horizontal='left' if j - 1 in left else 'right',
                                    vertical='center', wrap_text=True)
            c.border = Border(bottom=gold)
        ws.row_dimensions[r].height = 30
        r += 1
        for i, row in enumerate(rows):
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v)
                c.font = font(bold=j == 1, color=NAVY if j == 1 else INK)
                c.alignment = Alignment(horizontal='left' if j - 1 in left else 'right',
                                        vertical='top', wrap_text=j - 1 in left and j > 1,
                                        indent=1 if (j > 1 and j - 1 in left) else 0)
                c.border = Border(bottom=hair)
                if fmt and isinstance(v, (int, float)):
                    c.number_format = fmt
                if i % 2:
                    c.fill = fill(BAND)
            r += 1
        return r + 1

    ws = sheet('Overview', True, (30, 40, 4, 22, 60))
    r = title(ws, e['name'], f'{e.get("region")}  |  {e.get("country")}  |  index '
              f'{e.get("index") or "n/a"}  |  fund {e.get("etf") or "n/a"}  |  data as at {as_of}')
    ws.cell(r, 1, e.get('about') or '').font = font(italic=True, color=NAVY)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 2
    r = table(ws, r, ['Measure', 'Value'], _snapshot(ctx), left=(0, 1))
    s, c, v = ip.reading(ctx)
    ws.cell(r, 1, 'Where the numbers point').font = font(size=12, bold=True, color=NAVY)
    r += 1
    rows = [('Supported by the data', x) for x in s] + [('Cuts against it', x) for x in c] + \
           [('Verify before acting', x) for x in v]
    for i, (k, t) in enumerate(rows):
        ws.cell(r, 1, k).font = font(bold=True, color=NAVY)
        ws.cell(r, 2, t).font = font()
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 15 * max(1, len(t) // 100 + 1)
        for j in range(1, 6):
            ws.cell(r, j).border = Border(bottom=hair)
            if i % 2:
                ws.cell(r, j).fill = fill(BAND)
        r += 1

    ws = sheet('Returns', widths=(18, 24, 30, 18))
    r = title(ws, 'Returns over time', 'Three years and longer are annualised. The index '
              'excludes dividends; the fund includes them.')
    head, rows = _returns_rows(ctx)
    table(ws, r, head, [[x[0]] + [None if y is None else y / 100 for y in x[1:]] for x in rows],
          fmt='0.0%;-0.0%')

    ws = sheet('Risk by window', widths=(10, 12, 12, 14, 12, 12, 9, 12))
    r = title(ws, 'Return, drawdown and beta by time window', 'Beta and correlation against '
              'the S&P 500; weekly returns for a year or longer and for non-US markets.')
    r = table(ws, r, ['Window', 'Change', 'S&P 500', 'Difference (pts)', 'Worst fall',
                      'Volatility', 'Beta', 'Correlation'], _window_rows(ctx), fmt='0.0#')
    if ctx.get('dd_full'):
        ws.cell(r, 1, rk.drawdown_sentence(ctx['dd_full'], where='in the full record')).font = font(color=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=8)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
        r += 3
    rb = _rebased(ctx)
    if rb is not None and len(rb) > 10:
        hs = sheet('Price history', widths=(14, 16, 16))
        hs['A1'] = 'Weekly, rebased to 100 five years ago'
        hs['A1'].font = font(size=12, bold=True, color=NAVY)
        heads = ['Week ending', e['name']] + (['S&P 500'] if 'bench' in rb else [])
        for j, h in enumerate(heads, 1):
            cc = hs.cell(3, j, h)
            cc.font, cc.fill = font(bold=True, size=9, color=WHITE), fill(NAVY)
        for i, (d, row) in enumerate(rb.iterrows(), 4):
            hs.cell(i, 1, d.date()).number_format = 'dd mmm yyyy'
            hs.cell(i, 2, round(float(row['level']), 2))
            if 'bench' in rb and pd.notna(row['bench']):
                hs.cell(i, 3, round(float(row['bench']), 2))
        last = 3 + len(rb)
        ch = LineChart()
        ch.title = f'{e["name"]} against the S&P 500, rebased to 100'
        ch.height, ch.width = 8.5, 24
        ch.y_axis.majorGridlines = None
        ch.add_data(Reference(hs, min_col=2, max_col=2 + ('bench' in rb), min_row=3, max_row=last),
                    titles_from_data=True)
        ch.set_categories(Reference(hs, min_col=1, min_row=4, max_row=last))
        ch.x_axis.number_format = 'mmm yy'
        ch.series[0].graphicalProperties.line.solidFill = NAVY
        if len(ch.series) > 1:
            ch.series[1].graphicalProperties.line.solidFill = GOLD
        wb['Risk by window'].add_chart(ch, f'A{r}')

    if fd['holdings'] or fd['sectors']:
        ws = sheet('Holdings', widths=(14, 44, 14, 4, 28, 14))
        r = title(ws, f'What {e.get("etf")} holds', 'Largest holdings and sector split, from Yahoo.')
        if fd['holdings']:
            table(ws, r, ['Ticker', 'Company', 'Weight'],
                  [[h['symbol'], h['name'], None if h['weight'] is None else h['weight'] / 100]
                   for h in fd['holdings']], fmt='0.0%', left=(0, 1))
        if fd['sectors']:
            for j, h in enumerate(['Sector', 'Weight'], 5):
                cc = ws.cell(r, j, h)
                cc.font, cc.fill = font(bold=True, size=9, color=WHITE), fill(NAVY)
            for i, x in enumerate(fd['sectors'], r + 1):
                ws.cell(i, 5, x['sector']).font = font(bold=True, color=NAVY)
                ws.cell(i, 6, x['weight'] / 100).number_format = '0.0%'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ====================================================================== POWERPOINT =====
def build_index_pptx(ctx):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    e, fd = ctx['entry'], ctx['fund']
    as_of = ctx['end'].strftime('%d %B %Y').lstrip('0')
    HEAD, BODY = 'Georgia', 'Calibri'
    rgb = RGBColor.from_string
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    X0, CW = 0.6, 12.133
    foot = f'MICS International   |   {e["name"]}   |   Data as at {as_of}'
    n = {'i': 1}

    def text(sl, x, y, w, h, t, size=12, color=INK, face=BODY, bold=False, anchor=None, align=None):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor:
            tf.vertical_anchor = anchor
        for i, ln in enumerate(t if isinstance(t, list) else [t]):
            ln, o = ln if isinstance(ln, tuple) else (ln, {})
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align:
                p.alignment = align
            if o.get('space'):
                p.space_before = Pt(o['space'])
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.name = Pt(o.get('size', size)), face
            r.font.bold = o.get('bold', bold)
            r.font.color.rgb = rgb(o.get('color', color))

    def content(title, sub=None):
        n['i'] += 1
        sl = prs.slides.add_slide(blank)
        text(sl, X0, 0.38, CW, 0.75, title, size=24, color=NAVY, face=HEAD, anchor=MSO_ANCHOR.BOTTOM)
        ln = sl.shapes.add_connector(1, Inches(X0), Inches(1.22), Inches(X0 + CW), Inches(1.22))
        ln.line.color.rgb, ln.line.width = rgb(LINE), Pt(0.75)
        if sub:
            text(sl, X0, 1.3, CW, 0.3, sub, size=11, color=NAVY)
        text(sl, X0, 7.0, 9, 0.3, foot, size=9, color=NAVY)
        text(sl, 12.1, 7.0, 0.63, 0.3, str(n['i']), size=9, color=NAVY, align=PP_ALIGN.RIGHT)
        return sl

    def table(sl, x, y, colw, rows, size=10.5, rowh=0.34, left=(0,)):
        shp = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                  Inches(sum(colw)), Inches(rowh * len(rows)))
        tblPr = shp._element.graphic.graphicData.tbl.tblPr
        tblPr.set('bandRow', '0')
        tblPr.set('firstRow', '0')
        sid = tblPr.find(qn('a:tableStyleId'))
        if sid is not None:
            sid.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
        t = shp.table
        for j, w in enumerate(colw):
            t.columns[j].width = Inches(w)
        for i, row in enumerate(rows):
            t.rows[i].height = Inches(rowh)
            for j, v in enumerate(row):
                cell = t.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]
                cell.text_frame.word_wrap = True
                p.alignment = PP_ALIGN.LEFT if j in left else PP_ALIGN.RIGHT
                r = p.add_run()
                r.text = '' if v is None else str(v)
                r.font.name, r.font.size = BODY, Pt(size)
                r.font.bold = i == 0 or j == 0
                r.font.color.rgb = rgb(WHITE if i == 0 else NAVY if j == 0 else INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(NAVY if i == 0 else BAND if i % 2 == 0 else WHITE)

    pct = lambda v: fnum(v, 'pct')

    # cover
    sl = prs.slides.add_slide(blank)
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = rgb(NAVY)
    text(sl, X0, 1.6, 11.8, 1.8, e['name'], size=40, color=WHITE, face=HEAD, anchor=MSO_ANCHOR.BOTTOM)
    rule = sl.shapes.add_connector(1, Inches(X0), Inches(3.75), Inches(X0 + 1.4), Inches(3.75))
    rule.line.color.rgb, rule.line.width = rgb(GOLD), Pt(2)
    text(sl, X0, 4.0, 11.8, 1.8, [
        (f'{e.get("region")}  |  {e.get("country")}  |  fund {e.get("etf") or "n/a"}',
         {'color': WHITE, 'size': 16}),
        (e.get('about') or '', {'color': 'C9D3E3', 'size': 13, 'space': 6}),
        (f'Index analysis, data as at {as_of}   |   Prepared by MICS International',
         {'color': 'C9D3E3', 'size': 12, 'space': 6})])

    # at a glance + reading
    sl = content(f'{e["name"]} at a glance')
    table(sl, X0, 1.6, [2.9, 2.6], [['Measure', 'Value']] + _snapshot(ctx), size=11, rowh=0.38,
          left=(0, 1))
    s, c, v = ip.reading(ctx)
    none = 'Nothing material from these figures.'
    y = 1.6
    for head_, items in (('Supported by the data', s), ('Cuts against it', c),
                         ('Verify before acting', v)):
        text(sl, 6.5, y, 6.2, 0.3, head_, size=13, color=NAVY, face=HEAD)
        y += 0.35
        lines = [(f'{x}', {'space': 3}) for x in (items or [none])]
        text(sl, 6.5, y, 6.2, 0.4 * len(lines) + 0.3, lines, size=10.5, color=INK)
        y += 0.42 * len(lines) + 0.35

    # returns
    head, rr = _returns_rows(ctx)
    if rr:
        sl = content('Returns over time', 'Three years and longer are annualised. The index '
                     'excludes dividends; the fund includes them and is in its own currency.')
        rows = [head] + [[x[0]] + [pct(y) for y in x[1:]] for x in rr]
        w = (CW - 2.0) / (len(head) - 1)
        table(sl, X0, 1.8, [2.0] + [w] * (len(head) - 1), rows, size=11, rowh=0.4)

    # chart + window table: month end points, so the date axis stays readable
    rb = _rebased(ctx, freq='ME')
    if rb is not None and len(rb) > 10:
        yrs = max(1, round((rb.index[-1] - rb.index[0]).days / 365.25))
        sl = content(f'{e["name"]} against the S&P 500 over {yrs} years' if 'bench' in rb
                     else f'{e["name"]} over {yrs} years', 'Month end levels, rebased to 100 at the start')
        cd = CategoryChartData()
        cd.categories = [d.date() for d in rb.index]
        cd.add_series(e['name'], [round(float(x), 2) for x in rb['level']])
        if 'bench' in rb:
            cd.add_series('S&P 500', [None if pd.isna(x) else round(float(x), 2) for x in rb['bench']])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(X0), Inches(1.7), Inches(CW),
                                 Inches(5.1), cd).chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
        for i, ser in enumerate(ch.series):
            ser.smooth = False
            ser.format.line.color.rgb = rgb(NAVY if i == 0 else GOLD)
            ser.format.line.width = Pt(2 if i == 0 else 1.5)
        ca, va = ch.category_axis, ch.value_axis
        ca.tick_labels.font.size = va.tick_labels.font.size = Pt(9)
        ca.tick_labels.number_format, ca.tick_labels.number_format_is_linked = 'mmm yy', False
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(LINE)
        cax, anchor = ca._element, ca._element.find(qn('c:noMultiLvlLbl'))
        for tag in ('c:tickLblSkip', 'c:tickMarkSkip'):
            el = cax.makeelement(qn(tag), {'val': str(max(1, len(rb) // 8))})
            anchor.addprevious(el) if anchor is not None else cax.append(el)
    wr = _window_rows(ctx)
    if wr:
        sl = content('Return, drawdown and beta by time window',
                     'Beta and correlation against the S&P 500')
        rows = [['Window', 'Change', 'S&P 500', 'Difference', 'Worst fall', 'Volatility',
                 'Beta', 'Correlation']] + [
            [x[0], pct(x[1]), pct(x[2]), '-' if x[3] is None else f'{x[3]:+.1f} pts', pct(x[4]),
             pct(x[5]), fnum(x[6]), fnum(x[7])] for x in wr]
        table(sl, X0, 1.75, [1.3] + [(CW - 1.3) / 7] * 7, rows, size=11, rowh=0.4)
        if ctx.get('dd_full'):
            text(sl, X0, 1.75 + 0.4 * len(rows) + 0.3, CW, 0.8,
                 rk.drawdown_sentence(ctx['dd_full'], where='in the full record'), size=11, color=NAVY)

    # holdings
    if fd['holdings'] or fd['sectors']:
        sl = content(f'What {e.get("etf")} holds',
                     (f'The ten largest holdings are {fd["top10_pct"]:.0f}% of the fund'
                      if fd.get('top10_pct') else None))
        if fd['holdings']:
            rows = [['Ticker', 'Company', 'Weight']] + [[h['symbol'], h['name'], pct(h['weight'])]
                                                      for h in fd['holdings'][:12]]
            table(sl, X0, 1.75, [1.3, 4.0, 1.2], rows, size=10.5, rowh=0.36, left=(0, 1))
        if fd['sectors']:
            rows = [['Sector', 'Weight']] + [[x['sector'], pct(x['weight'])] for x in fd['sectors'][:12]]
            table(sl, 7.6, 1.75, [3.4, 1.73], rows, size=10.5, rowh=0.36)

    # notes
    sl = content('Notes and sources')
    table(sl, X0, 1.5, [2.2, CW - 2.2], [
        ['Item', 'Note'],
        ['Index', f'Yahoo, {e.get("index") or "not available, the fund is used instead"}. '
                  f'Price only, in {e.get("ccy") or "local currency"}.'],
        ['Fund', f'Yahoo, {e.get("etf") or "none"}. Adjusted for dividends, in the fund\'s own '
                 f'currency; holdings, costs and valuation as Yahoo reports them.'],
        ['S&P 500', 'Benchmark for beta, correlation and relative return, in US dollars.'],
        ['Use', 'For discussion and internal research. Not investment advice or a '
                'recommendation.']], size=11.5, rowh=0.6, left=(0, 1))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
